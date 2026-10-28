"""The workbook (V4 design doc §11): V4.3 writes the Model tab from a built model.

Rows are the model's lines, grouped by statement; columns its quarters and FY columns. Inputs
are values in blue with their source id as a cell note; every other cell is a live formula in
black, written from the same DSL expression Python evaluated (research/model/dsl.py), so the
sheet recomputes what Python did. tests/test_model_workbook.py has LibreOffice recalculate it
and compares every formula cell (the parity test). A cell Python leaves blank (an input
missing) is left blank, not written as a formula Excel would compute from empty cells as 0.
V4.4 adds forecast columns and the Assumptions tab (one blue value per scalar assumption, which
forecast formulas reference). The other tabs (Summary, Bridge, …) come with V4.5–V4.6.
"""

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tenk_agent.research.model.dsl import Cells, to_excel
from tenk_agent.research.model.graph import Model, expr_for

BLUE, BLACK = Font(color="0000FF"), Font(color="000000")
HEADER = Font(bold=True)
SECTION = PatternFill("solid", fgColor="D9DEE8")
FORECAST = PatternFill("solid", fgColor="F3F6FB")  # forecast columns, shaded
STATEMENTS = (("operating", "Operating build"), ("income", "Income statement"),
              ("cash_flow", "Cash flow"), ("balance", "Balance sheet"))
FORMATS = {"fraction": "0.0%", "USD per share": "0.00", "USD": "#,##0.00"}
FIRST_COLUMN = 3  # A: line, B: unit


def write(model: Model, path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    rows = layout(model, sheet)
    column_of = {c: FIRST_COLUMN + i for i, c in enumerate(model.columns)}
    cells = _cells(model, rows, column_of, assumptions_tab(model, book))
    for line in model.lines:
        for column in model.columns:
            cell = model.cells[(line.name, column)]
            target = sheet.cell(rows[line.name], column_of[column])
            target.number_format = FORMATS.get(line.unit, "#,##0")
            if cell.kind in ("input", "default"):
                target.value, target.font = cell.value, BLUE
                note = cell.ref or "default: not tagged by the company"
                target.comment = Comment(f"{note}\nfiled {cell.filed}" if cell.filed else note,
                                         "model")
            elif cell.value is not None:
                target.value, target.font = formula(model, line, column, cells), BLACK
                if column in model.forecast:
                    target.fill = FORECAST
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def layout(model: Model, sheet) -> dict[str, int]:
    """Header and statement sections; returns each line's row."""
    sheet.cell(1, 1, f"{model.ticker} model (history) as of {model.as_of}").font = HEADER
    sheet.cell(2, 1, "line").font = HEADER
    sheet.cell(2, 2, "unit").font = HEADER
    for i, column in enumerate(model.columns):
        sheet.cell(2, FIRST_COLUMN + i, column).font = HEADER
    sheet.column_dimensions["A"].width = 44
    rows, row = {}, 3
    for key, title in STATEMENTS:
        lines = [x for x in model.lines if x.statement == key]
        if not lines:
            continue
        sheet.cell(row, 1, title).font = HEADER
        for col in range(1, FIRST_COLUMN + len(model.columns)):
            sheet.cell(row, col).fill = SECTION
        row += 1
        for line in lines:
            sheet.cell(row, 1, line.label)
            sheet.cell(row, 2, line.unit)
            rows[line.name] = row
            row += 1
    return rows


def assumptions_tab(model: Model, book) -> dict[str, str]:
    """The Assumptions tab: name and value per assumption; returns each one's cell."""
    if not model.assumptions:
        return {}
    sheet = book.create_sheet("Assumptions")
    sheet.cell(1, 1, "assumption").font = HEADER
    sheet.cell(1, 2, "value").font = HEADER
    sheet.column_dimensions["A"].width = 36
    refs = {}
    for row, (name, value) in enumerate(sorted(model.assumptions.items()), start=2):
        sheet.cell(row, 1, name)
        cell = sheet.cell(row, 2, value)
        cell.font, cell.number_format = BLUE, "0.0000"
        refs[name] = f"Assumptions!$B${row}"
    return refs


def _cells(model: Model, rows: dict[str, int], column_of: dict[str, int],
           assumptions: dict[str, str] | None = None):
    def ref(name: str, column: str) -> str:
        return f"{get_column_letter(column_of[column])}{rows[name]}"

    def year_range(name: str, column: str) -> str | None:
        year = column[2:] if column.startswith("FY") else column[:4]
        quarters = [f"{year}Q{i}" for i in (1, 2, 3, 4)]
        if not all(q in column_of for q in quarters):
            return None
        return f"{ref(name, quarters[0])}:{ref(name, quarters[-1])}"

    return ref, year_range, assumptions or {}


def formula(model: Model, line, column: str, cells) -> str:
    """The cell's formula: an FY rule or the expression Python evaluated (graph.expr_for)."""
    ref, year_range, assumptions = cells
    expr = expr_for(model, line, column)
    if expr == "last":
        return f"={ref(line.name, f'{column[2:]}Q4')}"
    if expr == "sum":
        return f"=SUM({year_range(line.name, column)})"
    index = None if column.startswith("FY") else model.quarters.index(column)

    def cell(name: str, lag: int) -> str | None:
        if lag == 0:
            return ref(name, column)
        if index is None or index - lag < 0:
            return None
        return ref(name, model.quarters[index - lag])

    return to_excel(expr, Cells(cell, lambda name: year_range(name, column),
                                lambda name: assumptions[name]))


def parity(model: Model, path: Path) -> tuple[int, list[str]]:
    """LibreOffice recalculates the workbook at `path`; returns (formula cells compared, cells
    that differ from Python). Needs `soffice` on PATH."""
    soffice = shutil.which("soffice")
    if soffice is None:
        raise RuntimeError("the parity check needs LibreOffice (soffice) on PATH")
    formulas = load_workbook(path)["Model"]
    rows = layout(model, Workbook().active)  # the rows write() used, by line name
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp,
                        str(path)], check=True, capture_output=True, timeout=300)
        values = load_workbook(Path(tmp) / path.name, data_only=True)["Model"]
    checked, differ = 0, []
    for line in model.lines:
        for i, column in enumerate(model.columns):
            cell = formulas.cell(rows[line.name], FIRST_COLUMN + i)
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            checked += 1
            got, want = values.cell(cell.row, cell.column).value, model.value(line.name, column)
            if got is None or want is None or abs(got - want) > 1e-9 * max(1.0, abs(want)):
                differ.append(f"{line.name} {column}: {cell.value} = {got}, Python {want}")
    return checked, differ
