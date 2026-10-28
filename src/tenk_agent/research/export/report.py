"""S11 export (V4 design doc §11): the report's workbook and the note's PDF.

The workbook is the Model tab and Assumptions tab (research/export/workbook.py, checked for
parity with LibreOffice) plus Summary, Scenarios, Valuation, Revisions, Checks and Log (the
evidence and exception rows). The PDF is the note's Markdown printed with headless Chrome.
"""

import html
import re
from pathlib import Path

from openpyxl import load_workbook

from tenk_agent.research.export import workbook as model_tab
from tenk_agent.research.forecast import scenario_summary, scenarios


def workbook(folder: Path, model, f, valuation: dict, checks, facts, exceptions, revisions,
             settings) -> tuple[Path, str]:
    path = model_tab.write(model, folder / "workbook.xlsx")
    checked, differ = model_tab.parity(model, path)
    book = load_workbook(path)

    def sheet(title: str, header: list[str], rows: list[list]):
        ws = book.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = model_tab.HEADER
        for row in rows:
            ws.append(row)
        ws.column_dimensions["A"].width = 40

    v = valuation
    sheet("Summary", ["item", "value"], [
        ["rating (draft)", v["rating"]], ["price", v["price"]], ["price date", v["price_date"]],
        ["12-month price target", v["price_target"]], ["return", v["return"]],
        ["market cap (USD millions)", v["market_cap"]], ["EV (USD millions)", v["ev"]],
        ["net debt (USD millions)", v["net_debt"]], ["diluted shares (millions)", v["shares"]],
        ["WACC", v["cost_of_capital"]["wacc"]], ["parity", f"{checked - len(differ)}/{checked}"]])
    summary = scenario_summary(scenarios(f, settings), settings["scenarios"]["weights"])
    rows = [[f"{name} {line}", year, value] for name, lines in summary.items()
            for line, years in lines.items() for year, value in years.items()]
    sheet("Scenarios", ["scenario and line", "year", "value"], rows)
    sheet("Valuation", ["item", "value"], [
        *[[f"peer 2027E EV/EBITDA {t}", m] for t, m in v["peer_multiples"]["peers"].items()],
        ["peer median", v["peer_multiples"]["median"]],
        *[[f"{k}", x] for k, x in v["cost_of_capital"].items()],
        *[[f"{n} {k}", x] for n, s in v["scenarios"].items() for k, x in s.items()],
        *[[f"DCF per share, WACC|g {k}", x] for k, x in v["dcf_grid"].items()]])
    sheet("Revisions", ["line", "old", "new", "type", "confidence"],
          [[r["label"], r["old"], r["new"], r["type"], r["confidence"]] for r in revisions])
    sheet("Checks", ["check", "status", "tested", "failures", "skipped"],
          [[c.name, c.status, c.tested, "; ".join(c.failures), len(c.skipped)] for c in checks])
    sheet("Log", ["id", "kind", "label or severity", "value or item", "source"],
          [[f"E-{x.id[2:]}", "evidence", x.label, x.text(), x.source] for x in facts.items]
          + [[x["id"], "exception", x["severity"], x["item"], ""] for x in exceptions])
    book.save(path)
    return path, f"{checked - len(differ)}/{checked}"


def pdf(markdown_path: Path, out: Path) -> Path:
    from tenk_agent.evaluation.render import _print_pdf

    _print_pdf(to_html(markdown_path.read_text(), markdown_path.stem), out)
    return out


STYLE = """@page { size: A4; margin: 14mm 12mm; }
body { font: 9.5pt/1.4 -apple-system, "Segoe UI", Helvetica, Arial, sans-serif; color: #1a1a1a; }
h1 { font-size: 15pt; }
h2 { font-size: 11.5pt; margin: 5mm 0 1.5mm; border-bottom: 1px solid #999; }
table { border-collapse: collapse; margin: 1mm 0 2mm; } td, th { padding: .6mm 2.5mm;
border-bottom: 1px solid #ddd; text-align: right; }
td:first-child, th:first-child { text-align: left; }
th { background: #eef0f5; } p { margin: 1mm 0; }"""


def to_html(markdown: str, title: str) -> str:
    """The note's small Markdown subset: headings, tables, bullet lines and paragraphs."""
    out, table = [], []

    def flush():
        if table:
            rows = [r for r in table if not re.fullmatch(r"\|(---\|)+", r)]
            cells = [[html.escape(c.strip()) for c in r.strip("|").split("|")] for r in rows]
            out.append("<table><tr>" + "".join(f"<th>{c}</th>" for c in cells[0]) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>"
                                 for r in cells[1:]) + "</table>")
            table.clear()

    for line in markdown.splitlines():
        if line.startswith("|"):
            table.append(line)
            continue
        flush()
        if line.startswith("## "):
            out.append(f"<h2>{html.escape(line[3:])}</h2>")
        elif line.startswith("# "):
            out.append(f"<h1>{html.escape(line[2:])}</h1>")
        elif line.strip():
            out.append(f"<p>{html.escape(line)}</p>")
    flush()
    return (f"<!doctype html><meta charset=utf-8><title>{html.escape(title)}</title>"
            f"<style>{STYLE}</style>{''.join(out)}")
