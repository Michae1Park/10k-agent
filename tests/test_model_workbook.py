"""The Model tab (research/export/workbook.py) and its parity test: LibreOffice recalculates the
workbook and every formula cell must equal Python's value."""

import shutil

import pytest
from openpyxl import load_workbook

from tenk_agent.research.export import workbook
from test_model_graph import model, templates  # noqa: F401  (fixture)


def test_inputs_are_blue_values_formulas_black(templates, tmp_path):  # noqa: F811
    sheet = load_workbook(workbook.write(model(), tmp_path / "m.xlsx"))["Model"]
    rows = {sheet.cell(r, 1).value: r for r in range(1, sheet.max_row + 1)}
    first = workbook.FIRST_COLUMN
    revenue = sheet.cell(rows["revenue"], first)
    assert revenue.value == 10 and revenue.font.color.rgb.endswith("0000FF")
    assert revenue.comment.text.startswith("id:2025Q1")
    take_rate = sheet.cell(rows["take rate"], first)
    assert take_rate.value == f"=C{rows['revenue']}/C{rows['gross bookings']}"
    assert sheet.cell(rows["gb growth"], first).value is None  # its lag reaches before the data
    gb = rows["gross bookings"]  # 2026Q1 is column H; four quarters back skips FY2025 (G)
    assert sheet.cell(rows["gb growth"], first + 5).value == f"=(H{gb}/C{gb})-1"
    revenue_row = rows["revenue"]
    assert sheet.cell(revenue_row, first + 4).value == f"=SUM(C{revenue_row}:F{revenue_row})"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="needs LibreOffice")
def test_parity_with_libreoffice(templates, tmp_path):  # noqa: F811
    m = model()
    checked, differ = workbook.parity(m, workbook.write(m, tmp_path / "m.xlsx"))
    assert differ == [] and checked == 21
